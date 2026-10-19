import os
import sys
from pathlib import Path

import jsonschema
import openpyxl
import pytest
import yaml

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__)), "app", "src"))

from src.modality_contribution import (  # noqa: E402
    ProposalError,
    build_example_sidecar,
    build_schema,
    check_proposal,
    errors_only,
    read_proposal,
    write_proposal_workbook,
)

EXAMPLE_YAML = Path(__file__).resolve().parents[1] / "examples" / "modality_proposal" / "bloodwork.yml"


def _proposal(**over) -> dict:
    p = read_proposal(EXAMPLE_YAML)
    p.update(over)
    return p


def _no_vocabulary() -> dict:
    p = read_proposal(EXAMPLE_YAML)
    p["types"] = []
    for c in p["columns"]:
        c["type"] = None
    return p


def _messages(problems) -> str:
    return " | ".join(f"{p.where}: {p.message}" for p in errors_only(problems))


def test_example_defines_and_uses_a_vocabulary() -> None:
    p = read_proposal(EXAMPLE_YAML)
    assert p["types"][0] == {"value": "CRP", "description": "C-reactive protein"}
    assert [t["value"] for t in p["types"]] == ["CRP", "glucose", "cortisol"]
    cols = {c["id"]: c for c in p["columns"]}
    assert cols["crp"]["type"] == "CRP"
    assert cols["sample_id"]["type"] is None
    assert errors_only(check_proposal(p)) == []


def test_plain_string_types_are_accepted(tmp_path: Path) -> None:
    data = yaml.safe_load(EXAMPLE_YAML.read_text(encoding="utf-8"))
    data["types"] = ["CRP", "glucose", "cortisol"]
    f = tmp_path / "p.yml"
    f.write_text(yaml.safe_dump(data), encoding="utf-8")
    assert read_proposal(f)["types"] == [
        {"value": "CRP", "description": ""},
        {"value": "glucose", "description": ""},
        {"value": "cortisol", "description": ""},
    ]


def test_types_must_be_a_list_and_not_yaml_booleans(tmp_path: Path) -> None:
    data = yaml.safe_load(EXAMPLE_YAML.read_text(encoding="utf-8"))
    data["types"] = "CRP"
    f = tmp_path / "p.yml"
    f.write_text(yaml.safe_dump(data), encoding="utf-8")
    with pytest.raises(ProposalError, match="list"):
        read_proposal(f)
    f.write_text(EXAMPLE_YAML.read_text(encoding="utf-8").replace("value: CRP", "value: yes"), encoding="utf-8")
    with pytest.raises(ProposalError, match="quotes"):
        read_proposal(f)


def test_workbook_round_trip_keeps_types_and_column_types(tmp_path: Path) -> None:
    p = read_proposal(EXAMPLE_YAML)
    path = tmp_path / "p.xlsx"
    write_proposal_workbook(p, path)
    assert read_proposal(path) == p


def test_workbook_without_types_sheet_reads_as_no_vocabulary(tmp_path: Path) -> None:
    p = _no_vocabulary()
    path = tmp_path / "p.xlsx"
    write_proposal_workbook(p, path)
    wb = openpyxl.load_workbook(path)
    del wb["Types"]
    wb.save(path)
    assert read_proposal(path) == p


def test_type_header_is_not_mistaken_for_datatype(tmp_path: Path) -> None:
    p = read_proposal(EXAMPLE_YAML)
    path = tmp_path / "p.xlsx"
    write_proposal_workbook(p, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Columns"]
    headers = [c.value for c in ws[1]]
    a, b = headers.index("DataType"), headers.index("Type")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        row[a].value, row[b].value = row[b].value, row[a].value  # 'Type' now left of 'DataType'
    wb.save(path)
    assert read_proposal(path) == p


def test_column_type_must_be_in_the_vocabulary_or_other() -> None:
    p = read_proposal(EXAMPLE_YAML)
    cols = {c["id"]: c for c in p["columns"]}
    cols["crp"]["type"] = "CRPP"
    msg = _messages(check_proposal(p))
    assert "CRPP" in msg and "other" in msg and "not in the vocabulary" in msg
    cols["crp"]["type"] = "other"
    assert errors_only(check_proposal(p)) == []


def test_column_type_without_any_listed_types_is_an_error() -> None:
    p = _no_vocabulary()
    p["columns"][3]["type"] = "CRP"
    assert "defines no types" in _messages(check_proposal(p))


def test_vocabulary_values_are_checked() -> None:
    types = [
        {"value": "CRP", "description": ""},
        {"value": "CRP", "description": ""},
        {"value": "Other", "description": ""},
        {"value": "bad value", "description": ""},
        {"value": "", "description": ""},
    ]
    p = _proposal(types=types)
    for c in p["columns"]:
        c["type"] = None
    msg = _messages(check_proposal(p))
    assert "duplicate type 'CRP'" in msg
    assert "'other' is added automatically" in msg
    assert "letters, digits, underscores and hyphens" in msg
    assert "Value is empty" in msg


def test_schema_restricts_type_to_the_vocabulary_plus_other() -> None:
    p = read_proposal(EXAMPLE_YAML)
    schema = build_schema(p)
    type_prop = schema["additionalProperties"]["properties"]["Type"]
    assert type_prop["enum"] == ["CRP", "glucose", "cortisol", "other"]
    assert "C-reactive protein" in type_prop["description"]

    sidecar = build_example_sidecar(p, today="2026-10-15")
    assert sidecar["crp"]["Type"] == "CRP"
    assert "Type" not in sidecar["sample_id"]
    validator = jsonschema.Draft7Validator(schema)
    validator.validate(sidecar)

    sidecar["crp"]["Type"] = "other"
    validator.validate(sidecar)
    sidecar["crp"]["Type"] = "CRPP"
    assert list(validator.iter_errors(sidecar))


def test_schema_has_no_type_property_without_a_vocabulary() -> None:
    p = _no_vocabulary()
    assert "Type" not in build_schema(p)["additionalProperties"]["properties"]
    assert "Type" not in build_example_sidecar(p, today="2026-10-15")["crp"]
