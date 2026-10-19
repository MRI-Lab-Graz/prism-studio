import json
import os
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__)), "app", "src"))

from src.modality_contribution import (  # noqa: E402
    Problem,
    ProposalError,
    errors_only,
    read_proposal,
    write_proposal_workbook,
)

EXAMPLE_DIR = Path(__file__).resolve().parents[1] / "examples" / "modality_proposal"
EXAMPLE_YAML = EXAMPLE_DIR / "bloodwork.yml"


def test_read_yaml_example_normalizes_columns() -> None:
    p = read_proposal(EXAMPLE_YAML)
    assert (p["modality"], p["suffix"], p["extensions"]) == ("bloodwork", "bloodchem", ["tsv"])
    assert [c["id"] for c in p["columns"]] == [
        "sample_id", "draw_time", "fasting", "crp", "glucose", "cortisol",
    ]
    assert p["columns"][3] == {
        "id": "crp",
        "description": "C-reactive protein",
        "units": "mg/L",
        "datatype": "float",
        "scale": "continuous",
        "type": "CRP",
        "min": 0.0,
        "max": None,
        "allowed_values": None,
    }
    fasting = p["columns"][2]
    assert fasting["scale"] == "boolean"
    assert p["columns"][0]["scale"] is None  # blank stays 'not stated'
    assert fasting["allowed_values"] == ["yes", "no"]
    assert fasting["datatype"] == "string"
    assert fasting["units"] == "n/a"


def test_json_proposal_equals_yaml_proposal(tmp_path: Path) -> None:
    as_json = tmp_path / "p.json"
    as_json.write_text(
        json.dumps(yaml.safe_load(EXAMPLE_YAML.read_text(encoding="utf-8"))),
        encoding="utf-8",
    )
    assert read_proposal(as_json) == read_proposal(EXAMPLE_YAML)


def test_extensions_string_is_split_and_normalized(tmp_path: Path) -> None:
    f = tmp_path / "p.yml"
    f.write_text(
        EXAMPLE_YAML.read_text(encoding="utf-8").replace(
            "extensions: [tsv]", 'extensions: ".TSV, tsv.gz"'
        ),
        encoding="utf-8",
    )
    assert read_proposal(f)["extensions"] == ["tsv", "tsv.gz"]


def test_yaml_boolean_allowed_values_are_rejected_with_a_hint(tmp_path: Path) -> None:
    f = tmp_path / "p.yml"
    f.write_text(
        EXAMPLE_YAML.read_text(encoding="utf-8").replace(
            'allowed_values: ["yes", "no"]', "allowed_values: [yes, no]"
        ),
        encoding="utf-8",
    )
    with pytest.raises(ProposalError, match="quotes"):
        read_proposal(f)


def test_workbook_round_trip(tmp_path: Path) -> None:
    p = read_proposal(EXAMPLE_YAML)
    path = tmp_path / "p.xlsx"
    write_proposal_workbook(p, path)
    assert read_proposal(path) == p


def test_workbook_skips_blank_rows_and_ignores_extra_sheets(tmp_path: Path) -> None:
    p = read_proposal(EXAMPLE_YAML)
    path = tmp_path / "p.xlsx"
    write_proposal_workbook(p, path)
    wb = openpyxl.load_workbook(path)
    wb["Columns"].insert_rows(3)  # blank row inside the table
    wb.create_sheet("Notes")["A1"] = "scratch"
    wb.save(path)
    assert read_proposal(path) == p


def test_workbook_without_columns_sheet_is_a_clear_error(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Modality"
    path = tmp_path / "p.xlsx"
    wb.save(path)
    with pytest.raises(ProposalError, match="Columns"):
        read_proposal(path)


def test_unsupported_file_type_is_a_clear_error(tmp_path: Path) -> None:
    f = tmp_path / "p.txt"
    f.write_text("x", encoding="utf-8")
    with pytest.raises(ProposalError, match="unsupported"):
        read_proposal(f)


def test_scale_is_lowercased_and_never_guessed(tmp_path: Path) -> None:
    f = tmp_path / "p.yml"
    f.write_text(
        EXAMPLE_YAML.read_text(encoding="utf-8").replace("scale: boolean", "scale: Boolean"),
        encoding="utf-8",
    )
    cols = {c["id"]: c for c in read_proposal(f)["columns"]}
    assert cols["fasting"]["scale"] == "boolean"
    assert cols["sample_id"]["scale"] is None


def test_workbook_scale_cell_is_a_dropdown(tmp_path: Path) -> None:
    path = tmp_path / "p.xlsx"
    write_proposal_workbook(read_proposal(EXAMPLE_YAML), path)
    validations = openpyxl.load_workbook(path)["Columns"].data_validations.dataValidation
    assert any("continuous" in (v.formula1 or "") for v in validations)


def test_problem_str_and_errors_only() -> None:
    err = Problem("error", "Modality.Name", "bad")
    warn = Problem("warning", "x", "hmm")
    assert str(err) == "ERROR: Modality.Name: bad"
    assert errors_only([err, warn]) == [err]


def test_non_utf8_yaml_is_a_proposal_error(tmp_path: Path) -> None:
    f = tmp_path / "bad.yml"
    f.write_bytes(b"\xff\xfe\x00bad")
    with pytest.raises(ProposalError, match="cannot read bad.yml"):
        read_proposal(f)


def _yaml_with(tmp_path: Path, old: str, new: str) -> Path:
    f = tmp_path / "p.yml"
    text = EXAMPLE_YAML.read_text(encoding="utf-8")
    assert old in text
    f.write_text(text.replace(old, new, 1), encoding="utf-8")
    return f


@pytest.mark.parametrize("value", ["5", "true"])
def test_non_list_extensions_are_a_proposal_error(tmp_path: Path, value) -> None:
    f = _yaml_with(tmp_path, "extensions: [tsv]", f"extensions: {value}")
    with pytest.raises(ProposalError, match="'extensions' must be a list"):
        read_proposal(f)


@pytest.mark.parametrize("key", ["min", "max"])
def test_non_numeric_min_and_max_are_proposal_errors(key) -> None:
    from src.modality_contribution import _normalize_column

    with pytest.raises(ProposalError, match="not a number"):
        _normalize_column({"id": "x", key: [1]})
    with pytest.raises(ProposalError, match="not a number"):
        _normalize_column({"id": "x", key: "high"})


def test_workbook_comma_decimal_min_still_reads(tmp_path: Path) -> None:
    from src.modality_contribution import _normalize_column

    assert _normalize_column({"id": "x", "min": "1,5"})["min"] == 1.5


def test_yaml_boolean_type_and_scalar_allowed_values_are_rejected(tmp_path: Path) -> None:
    from src.modality_contribution import _normalize_column

    for key in ("type", "allowed_values"):
        with pytest.raises(ProposalError, match="quotes"):
            _normalize_column({"id": "x", key: True})
