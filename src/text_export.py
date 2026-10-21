"""Convert the Markdown subset our generators emit (methods, preregistration) to other formats.

Subset: ``## ``/``### `` headings, ``- `` bullets, paragraphs, ``**bold**``, ``*italic*``.
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

FORMATS = ("md", "html", "docx", "tex")
MIME = {
    "md": "text/markdown",
    "html": "text/html",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "tex": "application/x-tex",
}

_TEX_SPECIALS = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}
_SYMBOLS = {
    c: r"\ensuremath{\%s}" % name
    for c, name in zip(
        "αβγδηλμρσχω≥≤≠±×",
        "alpha beta gamma delta eta lambda mu rho sigma chi omega geq leq neq pm times".split(),
    )
}
_SYMBOLS["−"] = "-"
_UNDERSCORE_ITALIC = re.compile(r"(?<!\w)_(.+?)_(?!\w)")
_INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")


def _blocks(md: str):
    """Yield (kind, text) with kind in h2, h3, li, p."""
    for line in md.split("\n"):
        s = line.strip()
        if not s:
            continue
        if s.startswith("### "):
            yield "h3", s[4:]
        elif s.startswith("## "):
            yield "h2", s[3:]
        elif s.startswith("- "):
            yield "li", s[2:]
        else:
            yield "p", s


def _tex_escape(text: str) -> str:
    return "".join(_TEX_SPECIALS.get(c) or _SYMBOLS.get(c, c) for c in text)


def _normalize(md: str) -> str:
    """Preregistration Markdown uses ``# title`` and ``_italic_``; fold them into the supported subset."""
    md = re.sub(r"^# ", "## ", md, flags=re.M)
    return _UNDERSCORE_ITALIC.sub(r"*\1*", md)


def _tex_inline(text: str, cites: dict[str, str] | None = None) -> str:
    if cites:  # known citations become \\citep{key}; the rest is escaped as usual
        pattern = "|".join(
            r"\(" + re.escape(c) + r"\)|" + re.escape(c) for c in sorted(cites, key=len, reverse=True)
        )
        parts = re.split(f"({pattern})", text)
        return "".join(
            r"\citep{" + cites[p[1:-1] if p[1:-1] in cites else p] + "}"
            if i % 2
            else _tex_inline(p)
            for i, p in enumerate(parts)
        )
    out, pos = [], 0
    for m in _INLINE.finditer(text):
        out.append(_tex_escape(text[pos : m.start()]))
        if m.group(1) is not None:
            out.append(r"\textbf{" + _tex_escape(m.group(1)) + "}")
        else:
            out.append(r"\emph{" + _tex_escape(m.group(2)) + "}")
        pos = m.end()
    out.append(_tex_escape(text[pos:]))
    return "".join(out)


def md_to_tex(md: str, cites: dict[str, str] | None = None) -> str:
    """Article source; with ``cites`` (citation text -> bib key) known citations become ``\\citep``."""
    body: list[str] = []
    in_list = False
    for kind, text in _blocks(md):
        if kind != "li" and in_list:
            body.append(r"\end{itemize}")
            in_list = False
        if kind == "li" and not in_list:
            body.append(r"\begin{itemize}")
            in_list = True
        inline = _tex_inline(text, cites)
        if kind == "h2":
            body.append(r"\section{" + inline + "}")
        elif kind == "h3":
            body.append(r"\subsection{" + inline + "}")
        elif kind == "li":
            body.append(r"\item " + inline)
        else:
            body.append(inline + "\n")
    if in_list:
        body.append(r"\end{itemize}")
    text = "\n".join(body)
    cited = "\\citep{" in text
    return (
        "\\documentclass{article}\n\\usepackage[utf8]{inputenc}\n"
        "\\usepackage[T1]{fontenc}\n"
        + ("\\usepackage{natbib}\n" if cited else "")
        + "\\begin{document}\n\n"
        + text
        + ("\n\n\\bibliographystyle{apalike}\n\\bibliography{references}" if cited else "")
        + "\n\n\\end{document}\n"
    )


def _add_runs(paragraph, text: str) -> None:
    pos = 0
    for m in _INLINE.finditer(text):
        paragraph.add_run(text[pos : m.start()])
        if m.group(1) is not None:
            paragraph.add_run(m.group(1)).bold = True
        else:
            paragraph.add_run(m.group(2)).italic = True
        pos = m.end()
    paragraph.add_run(text[pos:])


def md_to_docx(md: str) -> bytes:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for Word export. "
            "Install it with: pip install python-docx"
        ) from exc
    doc = docx.Document()
    for kind, text in _blocks(md):
        if kind in ("h2", "h3"):
            doc.add_heading(text, level=2 if kind == "h2" else 3)
        else:
            _add_runs(doc.add_paragraph(style="List Bullet" if kind == "li" else None), text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def convert(md: str, fmt: str, cites: dict[str, str] | None = None) -> bytes:
    if fmt == "md":
        return md.encode("utf-8")
    md = _normalize(md)
    if fmt == "html":
        from src.reporting import _md_to_html

        return _md_to_html(md).encode("utf-8")
    if fmt == "docx":
        return md_to_docx(md)
    if fmt == "tex":
        return md_to_tex(md, cites).encode("utf-8")
    raise ValueError(f"Unknown format {fmt!r}; use one of {', '.join(FORMATS)}")


def _bib_escape(text: str) -> str:
    return "".join(r"\textbackslash{}" if c == "\\" else ("\\" + c if c in "&%$#_{}" else c) for c in text)


_APA_AUTHOR_YEAR = re.compile(r"^(.+?)\s*\((\d{4}[a-z]?)\)")


def _bib_entries(template_data: dict[str, dict]) -> list[tuple[str, str, dict]]:
    """``(key, citation text, reference)`` for the distinct references, in stable order.

    Keys are surname+year from the APA text (``Surname, X. (Year). ...``); colliding keys
    get ``a``, ``b`` ... in citation order, so re-export is identical.
    """
    refs: dict[str, dict] = {}  # citation text -> reference
    for tpl in template_data.values():
        study = tpl.get("Study") if isinstance(tpl, dict) else None
        for ref in (study or {}).get("References") or []:
            cit = " ".join(str(ref.get("Citation") or "").split()) if isinstance(ref, dict) else ""
            if cit:
                refs.setdefault(cit, ref)
    bases = {}
    for cit in refs:
        surname = re.sub(r"[^a-z]", "", cit.split(",")[0].split()[-1].lower()) if "," in cit else ""
        year = re.search(r"\((\d{4})", cit) or re.search(r"\b(\d{4})\b", cit)
        bases[cit] = (surname or "ref") + (year.group(1) if year else "")
    counts = {b: list(bases.values()).count(b) for b in bases.values()}
    seen: dict[str, int] = {}
    entries = []
    for cit in sorted(refs):
        key = bases[cit]
        if counts[key] > 1:
            key += chr(ord("a") + seen.get(bases[cit], 0))
            seen[bases[cit]] = seen.get(bases[cit], 0) + 1
        entries.append((key, cit, refs[cit]))
    return entries


def citation_keys(template_data: dict[str, dict]) -> dict[str, str]:
    """Citation text -> bib key, only for citations whose author and year could be read."""
    return {cit: key for key, cit, _ in _bib_entries(template_data) if _APA_AUTHOR_YEAR.match(cit)}


PRISM_BIB = """@misc{koschutnig2026prism,
  author = {{Koschutnig, K. and Weber, B. and Matischek, D.}},
  year = {2026},
  title = {{PRISM Studio}},
  doi = {10.5281/zenodo.22809100},
  url = {https://doi.org/10.5281/zenodo.22809100}
}
"""


def references_bib(template_data: dict[str, dict]) -> str:
    """BibTeX ``@misc`` entries for the distinct ``Study.References`` of all templates, plus PRISM Studio."""
    out = [PRISM_BIB]
    for key, cit, ref in _bib_entries(template_data):
        fields = []
        m = _APA_AUTHOR_YEAR.match(cit)
        if m:  # double braces: the APA author string is printed as written
            fields += [f"  author = {{{{{_bib_escape(m.group(1))}}}}}", f"  year = {{{m.group(2)}}}"]
        fields.append(f"  note = {{{_bib_escape(cit)}}}")
        for name, src in (("doi", "DOI"), ("url", "URL")):
            if ref.get(src):
                fields.append(f"  {name} = {{{ref[src]}}}")
        if not m and ref.get("Year"):
            fields.append(f"  year = {{{ref['Year']}}}")
        out.append(f"@misc{{{key},\n" + ",\n".join(fields) + "\n}\n")
    return "\n".join(out)


def emit_document(markdown, fmt: str, output: str | None, render=convert) -> None:
    """``render(markdown, fmt)`` to bytes, then print or write; exit 2 on a user error."""
    if fmt == "docx" and not output:
        print("Error: --format docx is binary and needs --output FILE")
        sys.exit(2)
    try:
        data = render(markdown, fmt)
    except RuntimeError as exc:  # python-docx missing
        print(f"Error: {exc}")
        sys.exit(2)
    if output:
        Path(output).write_bytes(data)
        print(f"Wrote {fmt} to {output}")
    else:
        print(data.decode("utf-8"))
